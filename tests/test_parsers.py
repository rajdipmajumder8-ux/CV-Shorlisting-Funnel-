import pytest
from pathlib import Path
import tempfile
import os

from src.parsers import extract_text, clean_text


class TestCleanText:
    def test_normalizes_whitespace(self):
        raw = "Hello   world\n\n\n\nMultiple   spaces"
        cleaned = clean_text(raw)
        assert "  " not in cleaned
        assert "\n\n\n" not in cleaned

    def test_strips_control_chars(self):
        raw = "Hello\x00world\x1f\x7f"
        cleaned = clean_text(raw)
        assert "\x00" not in cleaned
        assert "\x1f" not in cleaned
        assert "\x7f" not in cleaned

    def test_preserves_newlines_and_tabs(self):
        raw = "Line 1\nLine 2\tTabbed"
        cleaned = clean_text(raw)
        assert "\n" in cleaned
        assert "\t" in cleaned

    def test_empty_string(self):
        assert clean_text("") == ""

    def test_only_whitespace(self):
        assert clean_text("   \n\n\t  ") == ""


class TestExtractText:
    def test_txt_utf8(self):
        with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False, encoding='utf-8') as f:
            f.write("UTF-8 content: résumé, café")
            path = Path(f.name)
        try:
            text = extract_text(path)
            assert "résumé" in text
            assert "café" in text
        finally:
            os.unlink(path)

    def test_txt_latin1_fallback(self):
        with tempfile.NamedTemporaryFile(mode='wb', suffix='.txt', delete=False) as f:
            f.write(b"Latin-1: r\xe9sum\xe9")
            path = Path(f.name)
        try:
            text = extract_text(path)
            assert "résumé" in text or "resume" in text.lower()
        finally:
            os.unlink(path)

    def test_pdf_plumber_extraction(self):
        pytest.importorskip("pdfplumber")
        pdf_content = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>\nendobj\n4 0 obj\n<< /Length 44 >>\nstream\nBT /F1 12 Tf 100 700 Td (Test PDF content) Tj ET\nendstream\nendobj\nxref\n0 5\n0000000000 65535 f \n0000000009 00000 n \n0000000058 00000 n \n0000000115 00000 n \n0000000206 00000 n \ntrailer\n<< /Size 5 /Root 1 0 R >>\nstartxref\n299\n%%EOF"
        with tempfile.NamedTemporaryFile(mode='wb', suffix='.pdf', delete=False) as f:
            f.write(pdf_content)
            path = Path(f.name)
        try:
            text = extract_text(path)
            assert isinstance(text, str)
        finally:
            os.unlink(path)

    def test_corrupt_pdf_returns_empty_and_logs(self, caplog):
        with tempfile.NamedTemporaryFile(mode='wb', suffix='.pdf', delete=False) as f:
            f.write(b"not a pdf at all")
            path = Path(f.name)
        try:
            text = extract_text(path)
            assert text == ""
            # Use caplog.text which aggregates all log output
            log_text = caplog.text.lower()
            assert "pdf" in log_text
            assert "fallback" in log_text or "error" in log_text
        finally:
            os.unlink(path)

    def test_docx_extraction(self):
        pytest.importorskip("docx")
        from docx import Document
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            doc = Document()
            doc.add_paragraph("Paragraph 1")
            doc.add_paragraph("Paragraph 2")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Cell A1"
            table.cell(0, 1).text = "Cell B1"
            table.cell(1, 0).text = "Cell A2"
            table.cell(1, 1).text = "Cell B2"
            doc.save(f.name)
            path = Path(f.name)
        try:
            text = extract_text(path)
            assert "Paragraph 1" in text
            assert "Paragraph 2" in text
            assert "Cell A1" in text
            assert "Cell B2" in text
        finally:
            os.unlink(path)

    def test_nonexistent_file_raises(self):
        with pytest.raises(FileNotFoundError):
            extract_text(Path("/nonexistent/file.pdf"))

    def test_unsupported_extension_raises(self):
        with tempfile.NamedTemporaryFile(suffix='.xyz', delete=False) as f:
            path = Path(f.name)
        try:
            with pytest.raises(ValueError, match="Unsupported file type"):
                extract_text(path)
        finally:
            os.unlink(path)