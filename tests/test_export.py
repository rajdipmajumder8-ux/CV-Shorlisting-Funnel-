import pytest
from io import BytesIO
from openpyxl import load_workbook

from src.export import build_excel_report
from src.scoring.tfidf_scorer import ScoreResult


class TestExcelExport:
    def test_build_excel_report_returns_bytesio(self):
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python", "Flask", "PostgreSQL"],
                missing_skills=["AWS"],
                pros=["Strong alignment", "Key skills: Python, Flask"],
                cons=["Missing: AWS"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            ),
            ScoreResult(
                candidate_name="Jane Smith",
                file_name="jane.pdf",
                score=0.60,
                tier=2,
                matched_skills=["Python", "Flask"],
                missing_skills=["PostgreSQL", "Docker"],
                pros=["Moderate alignment", "Relevant: Python, Flask"],
                cons=["Missing: PostgreSQL, Docker"],
                next_action="Technical screen recommended",
                raw_breakdown={"method": "tfidf", "similarity": 0.60}
            ),
            ScoreResult(
                candidate_name="Bob Wilson",
                file_name="bob.pdf",
                score=0.30,
                tier=3,
                matched_skills=[],
                missing_skills=["Python", "Flask", "PostgreSQL"],
                pros=["Limited overlap"],
                cons=["Significant gaps: Python, Flask, PostgreSQL"],
                next_action="Consider for future roles / reject",
                raw_breakdown={"method": "tfidf", "similarity": 0.30}
            )
        ]
        
        bio = build_excel_report(results)
        
        assert isinstance(bio, BytesIO)
        assert bio.tell() == 0
        assert len(bio.getvalue()) > 0

    def test_excel_has_summary_and_detail_sheets(self):
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python"],
                missing_skills=["AWS"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            )
        ]
        
        bio = build_excel_report(results)
        wb = load_workbook(bio)
        
        # Should have Summary Matrix + 1 detail sheet per candidate
        assert "Summary Matrix" in wb.sheetnames
        assert "John Doe" in wb.sheetnames
        assert len(wb.sheetnames) == 2

    def test_excel_multiple_candidates_multiple_detail_sheets(self):
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python"],
                missing_skills=["AWS"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            ),
            ScoreResult(
                candidate_name="Jane Smith",
                file_name="jane.pdf",
                score=0.60,
                tier=2,
                matched_skills=["Python"],
                missing_skills=["Docker"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Technical screen recommended",
                raw_breakdown={"method": "tfidf", "similarity": 0.60}
            ),
            ScoreResult(
                candidate_name="Bob Wilson",
                file_name="bob.pdf",
                score=0.30,
                tier=3,
                matched_skills=[],
                missing_skills=["Python"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Consider for future roles / reject",
                raw_breakdown={"method": "tfidf", "similarity": 0.30}
            )
        ]
        
        bio = build_excel_report(results)
        wb = load_workbook(bio)
        
        assert "Summary Matrix" in wb.sheetnames
        assert "John Doe" in wb.sheetnames
        assert "Jane Smith" in wb.sheetnames
        assert "Bob Wilson" in wb.sheetnames
        assert len(wb.sheetnames) == 4

    def test_summary_sheet_has_title_and_headers(self):
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python", "Flask"],
                missing_skills=["AWS"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            )
        ]
        
        bio = build_excel_report(results)
        wb = load_workbook(bio)
        ws = wb["Summary Matrix"]
        
        # Title row (row 1)
        assert "Sunjet Energy" in str(ws.cell(row=1, column=1).value)
        assert "Talent Funnel" in str(ws.cell(row=1, column=1).value)
        
        # Subtitle row (row 2)
        assert "Generated:" in str(ws.cell(row=2, column=1).value)
        
        # Headers at row 4
        headers = [ws.cell(row=4, column=c).value for c in range(1, 8)]
        expected = ["Candidate", "File", "Score %", "Tier", "Key Matched Skills", "Missing Skills", "Next Action"]
        assert headers == expected

    def test_summary_sheet_header_styling(self):
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python"],
                missing_skills=["AWS"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            )
        ]
        
        bio = build_excel_report(results)
        wb = load_workbook(bio)
        ws = wb["Summary Matrix"]
        
        # Check header styling at row 4
        for col in range(1, 8):
            cell = ws.cell(row=4, column=col)
            assert cell.fill.start_color.rgb == "001B2A4A"  # Dark blue
            assert cell.font.color.rgb == "00FFFFFF"  # White
            assert cell.font.bold is True

    def test_tier_conditional_formatting(self):
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python"],
                missing_skills=["AWS"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            ),
            ScoreResult(
                candidate_name="Jane Smith",
                file_name="jane.pdf",
                score=0.60,
                tier=2,
                matched_skills=["Python"],
                missing_skills=["Docker"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Technical screen recommended",
                raw_breakdown={"method": "tfidf", "similarity": 0.60}
            ),
            ScoreResult(
                candidate_name="Bob Wilson",
                file_name="bob.pdf",
                score=0.30,
                tier=3,
                matched_skills=[],
                missing_skills=["Python"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Consider for future roles / reject",
                raw_breakdown={"method": "tfidf", "similarity": 0.30}
            )
        ]
        
        bio = build_excel_report(results)
        wb = load_workbook(bio)
        ws = wb["Summary Matrix"]
        
        # Data starts at row 5 (after title, subtitle, blank, headers)
        # Check tier column (column 4) has tier colors
        tier1_cell = ws.cell(row=5, column=4)  # John Doe - Tier 1
        tier2_cell = ws.cell(row=6, column=4)  # Jane Smith - Tier 2
        tier3_cell = ws.cell(row=7, column=4)  # Bob Wilson - Tier 3
        
        assert tier1_cell.fill.start_color.rgb == "00C6EFCE"
        assert tier2_cell.fill.start_color.rgb == "00FFEB9C"
        assert tier3_cell.fill.start_color.rgb == "00FFC7CE"

    def test_detail_sheet_has_professional_sections(self):
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python", "Flask"],
                missing_skills=["AWS", "Docker"],
                pros=["Strong alignment", "Key skills: Python, Flask"],
                cons=["Missing: AWS, Docker"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            )
        ]
        
        bio = build_excel_report(results)
        wb = load_workbook(bio)
        ws = wb["John Doe"]
        
        # Check sections exist
        all_values = []
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell:
                    all_values.append(str(cell))
        
        # New professional section headers
        assert "Candidate Profile: John Doe" in all_values
        assert "Matched Skills" in all_values or "✓ Matched Skills" in " ".join(all_values)
        assert "Missing" in " ".join(all_values) or "Gap Skills" in " ".join(all_values)
        assert "Assessment Rationale" in all_values
        assert "Strengths" in " ".join(all_values) or "Pros" in all_values
        assert "Improvement" in " ".join(all_values) or "Cons" in all_values
        assert "Next Action" in all_values or "Recommended" in " ".join(all_values)
        assert "Python" in all_values
        assert "Flask" in all_values
        assert "AWS" in all_values
        assert "Schedule interview" in all_values

    def test_frozen_pane_and_autofilter(self):
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python"],
                missing_skills=["AWS"],
                pros=["Pro 1"],
                cons=["Con 1"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            )
        ]
        
        bio = build_excel_report(results)
        wb = load_workbook(bio)
        ws = wb["Summary Matrix"]
        
        # Frozen pane at row 5 (after title, subtitle, blank, headers)
        assert ws.freeze_panes == "A5"
        
        # Auto-filter
        assert ws.auto_filter.ref is not None

    def test_missing_skills_readable_format(self):
        """Test that missing skills are formatted with line breaks for readability."""
        results = [
            ScoreResult(
                candidate_name="John Doe",
                file_name="john.pdf",
                score=0.85,
                tier=1,
                matched_skills=["Python", "Flask", "Django", "PostgreSQL", "Docker", "AWS"],
                missing_skills=["Kubernetes", "GCP", "CI/CD", "Terraform", "GraphQL"],
                pros=["Strong alignment"],
                cons=["Missing cloud skills"],
                next_action="Schedule interview",
                raw_breakdown={"method": "tfidf", "similarity": 0.85}
            )
        ]
        
        bio = build_excel_report(results)
        wb = load_workbook(bio)
        ws = wb["Summary Matrix"]
        
        # Missing skills cell should have line breaks (not just comma separated)
        missing_cell = ws.cell(row=5, column=6)  # Column 6 = Missing Skills
        missing_text = str(missing_cell.value)
        assert "\n" in missing_text or ";" in missing_text
        # Should contain the skills
        assert "Kubernetes" in missing_text
        assert "GCP" in missing_text